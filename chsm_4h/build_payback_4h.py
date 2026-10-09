# -*- coding: utf-8 -*-
"""Окупаемость разгрузки 4 часов по заводам при точности попадания в ЧСМ 80 %.

Сценарий 1 («как задано»): завод снижает нагрузку на свой потолок в 4 часах,
ЧСМ попадает в эти 4 часа в 80 % рабочих дней.
    зачётное снижение = потолок × точность
    экономия в месяц  = зачётное снижение × цена мощности
    окупаемость       = инвестиции ÷ экономия

Сценарий 2 («поправка на перенос пика»): то же, но срез в каждом попадании
не больше запаса блока над остальными часами окна СО ЕЭС, иначе разгрузка
сама переносит ЧСМ за пределы блока. Доля безопасного среза считается
по посуточному факту АТС (идеальный посуточный выбор блока), то есть
это оценка сверху.

ЧСМ ищется только внутри плановых пиковых часов СО ЕЭС, блок — 4 часа
подряд целиком внутри окна.

Запуск:
    python build_payback_4h.py <папка raw с отчётами АТС> [выходной.xlsx]
"""
from __future__ import annotations

import statistics
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import accuracy as A  # noqa: E402
import load_ats  # noqa: E402

ACCURACY = 0.80
HOURS = A.BLOCK

#: Плановые часы пиковой нагрузки СО ЕЭС, МСК, приказ на 2026 г.
#: (для 2025 г. в so_data.py окна те же).
SO_WINDOWS = {
    1: {1: "8-21", 2: "8-13,17-21", 3: "8-21", 4: "8-15,18-21", 5: "8-15,20-21",
        6: "8-16,20-21", 7: "8-17,20-21", 8: "8-21", 9: "8-15,18-21",
        10: "8-21", 11: "8-11,16-21", 12: "8-12,15-21"},
    2: {1: "5-8,11-17", 2: "5-8,12-17", 3: "5-8,13-17", 4: "5-8,13-17",
        5: "5-8,13-17", 6: "5-11,13-17", 7: "5-17", 8: "5-10,14-17",
        9: "5-8,13-17", 10: "5-8,12-17", 11: "5-8,11-17", 12: "5-8,11-17"},
}

#: завод, город, регион, ценовая зона, потолок снижения МВт, примечание.
#: Потолки — zavody_data.MAX_DEPTH; КрАЗ — глубина модуляции из so_data.PLANTS.
PLANTS = [
    ("БрАЗ", "Братск", "Иркутская область", 2, 49.0, ""),
    ("ИркАЗ", "Шелехов", "Иркутская область", 2, 10.5, ""),
    ("ТАЗ", "Тайшет", "Иркутская область", 2, 6.0, ""),
    ("НкАЗ", "Новокузнецк", "Кемеровская область", 2, 17.0, ""),
    ("САЗ", "Саяногорск", "Республика Хакасия", 2, 15.0, ""),
    ("ВгАЗ", "Волгоград", "Волгоградская область", 1, 3.0, ""),
    ("КАЗ", "Кандалакша", "Мурманская область", 1, 1.0,
     "нет посуточного факта АТС"),
    ("КрАЗ", "Красноярск", "Красноярский край", 2, 49.0,
     "потолок из so_data (в реестре zavody_data нет)"),
]

INVEST = 1_000_000  # руб на завод, предварительно (как в zavody_data.INVEST)

#: Цена мощности, руб/МВт·мес — zavody_data.PRICES.
PRICES = [(2026, 1_000_000), (2027, 1_330_000), (2028, 1_670_000),
          (2029, 2_000_000), (2030, 2_330_000), (2031, 2_670_000),
          (2032, 3_000_000), (2033, 3_330_000), (2034, 3_670_000),
          (2035, 4_000_000)]

MONTH_RU = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль",
            "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]

FONT = "Times New Roman"
INK = "1F3864"
HDR = PatternFill("solid", fgColor=INK)
INP = PatternFill("solid", fgColor="FFF2CC")
TOT = PatternFill("solid", fgColor="DDEBF7")
OK = PatternFill("solid", fgColor="E2EFDA")
BAD = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
F_T = Font(name=FONT, bold=True, size=14, color=INK)
F_S = Font(name=FONT, size=11, color="404040")
F_H = Font(name=FONT, bold=True, size=12, color=INK)
F_HD = Font(name=FONT, bold=True, size=10, color="FFFFFF")
F = Font(name=FONT, size=11)
F_B = Font(name=FONT, bold=True, size=11)
F_SM = Font(name=FONT, size=9, color="595959")

RUB, MW, PCT = "#,##0", "0.00", "0%"


def window(ym: str, zone: int) -> set[int]:
    return A.parse_window(SO_WINDOWS[zone][int(ym[5:7])])


def ru_month(ym: str) -> str:
    return f"{MONTH_RU[int(ym[5:7]) - 1]} {ym[:4]}"


def head(ws, row, titles, height=40, col0=1):
    for j, t in enumerate(titles, start=col0):
        c = ws.cell(row, j, t)
        c.fill, c.font, c.border, c.alignment = HDR, F_HD, BOX, WRAP
    ws.row_dimensions[row].height = height


def cell(ws, r, c, v, fmt=None, font=F, fill=None, align=CENTER):
    x = ws.cell(r, c, v)
    x.font, x.border, x.alignment = font, BOX, align
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    return x


def title(ws, name, sub):
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, name).font = F_T
    ws.cell(2, 1, sub).font = F_S


# ------------------------------------------------------------------ расчёт
def region_facts(data):
    """Факты по регионам с заводами: блоки, попадания, безопасный срез."""
    out = {}
    for region in dict.fromkeys(p[2] for p in PLANTS):
        if region not in data:
            continue
        zone = next(p[3] for p in PLANTS if p[2] == region)
        depth = sum(p[4] for p in PLANTS if p[2] == region)
        wf = lambda ym, z=zone: window(ym, z)  # noqa: E731
        rows = A.analyse(data[region], wf, depth)
        orc = A.oracle(data[region], wf, depth)
        n = sum(r["n"] for r in rows)
        out[region] = dict(
            zone=zone, depth=depth, rows=rows, oracle=orc, n=n,
            plants=", ".join(p[0] for p in PLANTS if p[2] == region),
            hit0=sum(r["hit0"] for r in rows) / n,
            hit1=sum(r["hit1"] for r in rows) / n,
            full=sum(o["full"] for o in orc.values()) / n,
            safe=sum(o["safe"] * o["n"] for o in orc.values()) / n,
        )
    return out


#: Пояс относительно МСК, в котором лежат monthly_profiles (raw — всегда МСК).
PROFILE_TZ = {"Красноярский край": 4, "Иркутская область": 5,
              "Кемеровская область": 4, "Республика Хакасия": 4}


def check_profiles(data, profiles_dir):
    """Сверка: средний профиль рабочих дней = monthly_profiles/<регион>/<мес>.xlsx.

    Профили Сибири там в местном времени, поэтому час сдвигается на пояс.
    """
    if not profiles_dir:
        return
    for region, days in data.items():
        tz, worst = PROFILE_TZ.get(region, 0), 0.0
        for f in sorted(Path(profiles_dir, region).glob("*.xlsx")):
            ws = openpyxl.load_workbook(f, data_only=True).active
            for row in ws.iter_rows(min_row=3, values_only=True):
                if not isinstance(row[0], (int, float)):
                    continue
                h_msk = (int(row[0]) - 1 - tz) % 24 + 1
                mean = statistics.mean(d[h_msk] for k, d in days.items() if k[:7] == f.stem)
                worst = max(worst, abs(mean - float(row[1])))
        print(f"Сверка с monthly_profiles, {region}: макс. расхождение {worst:.1f} МВт")


# ------------------------------------------------------------------ листы
def build_svod(ws, facts):
    title(ws, "Окупаемость разгрузки 4 часов по заводам",
          "ЧСМ — только внутри плановых пиковых часов СО ЕЭС. "
          "Жёлтые ячейки — ввод, остальное — формулы.")

    ws.cell(4, 1, "Исходные данные").font = F_H
    cell(ws, 5, 1, "Точность попадания ЧСМ в 4 часа разгрузки", align=LEFT)
    acc = cell(ws, 5, 4, ACCURACY, PCT, F_B, INP)
    cell(ws, 6, 1, "Часов разгрузки в сутки", align=LEFT)
    cell(ws, 6, 4, HOURS, "0", F_B, INP)
    cell(ws, 7, 1, "Год цены для окупаемости", align=LEFT)
    cell(ws, 7, 4, 2026, "0", F_B, INP)
    for r in (5, 6, 7):
        for c in (2, 3):
            ws.cell(r, c).border = BOX
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ACC, YEAR = "$D$5", "$D$7"

    pr = 9
    ws.cell(pr, 1, "Цена мощности, руб/МВт·мес").font = F_H
    cell(ws, pr + 1, 1, "Год", font=F_HD, fill=HDR)
    cell(ws, pr + 2, 1, "Цена", font=F_HD, fill=HDR)
    for k, (y, p) in enumerate(PRICES):
        cell(ws, pr + 1, 2 + k, y, "0", F_HD, HDR)
        cell(ws, pr + 2, 2 + k, p, RUB, F_B, INP)
    last_pc = get_column_letter(1 + len(PRICES))
    YRS = f"$B${pr + 1}:${last_pc}${pr + 1}"
    PRC = f"$B${pr + 2}:${last_pc}${pr + 2}"
    PRICE = f"INDEX({PRC},MATCH({YEAR},{YRS},0))"

    hr = pr + 5
    ws.cell(hr - 1, 1, "Расчёт по заводам").font = F_H
    head(ws, hr, [
        "№", "Завод", "Регион", "Зона", "Потолок\nснижения, МВт",
        "Инвестиции,\nруб",
        # сценарий 1
        "Зачётное\nснижение, МВт", "Экономия\nв месяц, руб", "Экономия\nв год, руб",
        "Окупаемость,\nдней",
        # сценарий 2
        "Доля безопасного\nсреза (факт АТС)", "Зачётное\nснижение, МВт",
        "Экономия\nв год, руб", "Окупаемость,\nдней",
        # горизонт
        "Экономия 2026–2035,\nруб (сц. 1)", "Экономия 2026–2035,\nруб (сц. 2)",
        "Примечание"], height=48)
    ws.cell(hr - 1, 7, "Сценарий 1: 80 % × потолок").font = F_B
    ws.cell(hr - 1, 11, "Сценарий 2: поправка на перенос пика").font = F_B

    r0 = hr + 1
    for i, (name, city, region, zone, depth, note) in enumerate(PLANTS):
        r = r0 + i
        fct = facts.get(region)
        cell(ws, r, 1, i + 1)
        cell(ws, r, 2, name, font=F_B, align=LEFT)
        cell(ws, r, 3, region, align=LEFT)
        cell(ws, r, 4, zone)
        cell(ws, r, 5, depth, MW, F_B, INP)
        cell(ws, r, 6, INVEST, RUB, F_B, INP)
        cell(ws, r, 7, f"=E{r}*{ACC}", MW)
        cell(ws, r, 8, f"=G{r}*{PRICE}", RUB)
        cell(ws, r, 9, f"=H{r}*12", RUB, F_B)
        cell(ws, r, 10, f'=IF(I{r}<=0,"—",F{r}/I{r}*365)', "0.0", F_B, OK)
        cell(ws, r, 11, round(fct["safe"] / fct["depth"], 4) if fct else "н/д", PCT)
        cell(ws, r, 12, f'=IF(ISNUMBER(K{r}),G{r}*K{r},"н/д")', MW)
        cell(ws, r, 13, f'=IF(ISNUMBER(L{r}),L{r}*{PRICE}*12,"н/д")', RUB, F_B)
        cell(ws, r, 14, f'=IF(ISNUMBER(M{r}),IF(M{r}<=0,"—",F{r}/M{r}*365),"н/д")',
             "0.0", F_B, OK)
        cell(ws, r, 15, f"=G{r}*12*SUM({PRC})", RUB)
        cell(ws, r, 16, f'=IF(ISNUMBER(L{r}),L{r}*12*SUM({PRC}),"н/д")', RUB)
        cell(ws, r, 17, note, font=F_SM, align=LEFT)
    tr = r0 + len(PLANTS)
    cell(ws, tr, 2, "ИТОГО", font=F_B, fill=TOT, align=LEFT)
    for c in range(1, 18):
        x = ws.cell(tr, c)
        x.border, x.fill, x.font = BOX, TOT, F_B
        x.alignment = CENTER
    for c, fmt in ((5, MW), (6, RUB), (7, MW), (8, RUB), (9, RUB), (12, MW), (13, RUB),
                   (15, RUB), (16, RUB)):
        L = get_column_letter(c)
        ws.cell(tr, c, f"=SUM({L}{r0}:{L}{tr - 1})").number_format = fmt
    ws.cell(tr, 10, f'=IF(I{tr}<=0,"—",F{tr}/I{tr}*365)').number_format = "0.0"
    ws.cell(tr, 14, f'=IF(M{tr}<=0,"—",F{tr}/M{tr}*365)').number_format = "0.0"

    notes = [
        "Сценарий 1 — как задано: потолок снижения держится все 4 часа, ЧСМ попадает "
        "в эти часы в 80 % рабочих дней; зачётное снижение = потолок × 80 %.",
        "Сценарий 2 — в дни попадания срез не больше запаса блока над остальными "
        "часами окна СО ЕЭС (иначе разгрузка сама уводит ЧСМ из блока). Доля "
        "безопасного среза — по посуточному факту АТС июнь 2025 — май 2026, для "
        "региона в целом (заводы региона режут совместно).",
        "Инвестиции 1 млн руб на завод — предварительное значение из zavody_data; "
        "при нём окупаемость измеряется днями. Подставьте реальные затраты в столбец F.",
        "Не учтены: недовыпуск/перенос производства в часы разгрузки, изменение "
        "стоимости электроэнергии, сетевая составляющая.",
    ]
    for k, t in enumerate(notes):
        ws.cell(tr + 2 + k, 1, t).font = F_SM

    for col, w in {"A": 5, "B": 9, "C": 23, "D": 6, "E": 12, "F": 13, "G": 11,
                   "H": 15, "I": 16, "J": 13, "K": 15, "L": 11, "M": 16, "N": 13,
                   "O": 19, "P": 19, "Q": 30}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"C{r0}"
    return r0, tr


def build_check(ws, facts):
    title(ws, "Проверка точности по факту АТС",
          "Рабочие дни июнь 2025 — май 2026. ЧСМ = час максимума потребления "
          "субъекта внутри окна СО ЕЭС. Блок — 4 часа подряд внутри окна.")
    hr = 4
    head(ws, hr, [
        "Регион", "Зона", "Заводы", "Суммарный\nсрез, МВт", "Раб.\nдней",
        "Фикс. блок на месяц:\nЧСМ в блоке\nбез разгрузки",
        "Фикс. блок на месяц:\nЧСМ в блоке\nпосле разгрузки",
        "Идеальный прогноз:\nдней, где полный срез\nне уводит ЧСМ",
        "Идеальный прогноз:\nсредний безопасный\nсрез, МВт",
        "Доля\nбезопасного\nсреза"], height=62)
    r = hr + 1
    for region, f in facts.items():
        cell(ws, r, 1, region, align=LEFT, font=F_B)
        cell(ws, r, 2, f["zone"])
        cell(ws, r, 3, f["plants"], align=LEFT)
        cell(ws, r, 4, f["depth"], MW)
        cell(ws, r, 5, f["n"])
        for c, v in ((6, f["hit0"]), (7, f["hit1"]), (8, f["full"])):
            cell(ws, r, c, round(v, 4), PCT, F_B, OK if v >= ACCURACY else BAD)
        cell(ws, r, 9, round(f["safe"], 2), MW)
        cell(ws, r, 10, round(f["safe"] / f["depth"], 4), PCT, F_B)
        r += 1
    ws.cell(r, 1, "Зелёный — не ниже 80 %, красный — ниже. Фиксированный блок "
                  "подобран задним числом по самому месяцу (оценка сверху для "
                  "расписания «один блок на месяц»).").font = F_SM
    ws.cell(r + 1, 1, "Идеальный прогноз — на каждые сутки выбран лучший блок; "
                      "это потолок для любой модели прогноза.").font = F_SM

    r += 3
    for region, f in facts.items():
        ws.cell(r, 1, f"{region} — по месяцам (срез {f['depth']:g} МВт)").font = F_H
        r += 1
        head(ws, r, ["Месяц", "Раб.\nдней", "Окно СО ЕЭС, МСК", "Лучший\nблок",
                     "ЧСМ в блоке\nбез разгрузки", "ЧСМ в блоке\nпосле разгрузки",
                     "Полный срез\nне уводит ЧСМ\n(идеал)", "Безопасный\nсрез, МВт\n(идеал)",
                     "Частые часы ЧСМ\n(час: дней)"], height=50)
        r += 1
        for row in f["rows"]:
            o = f["oracle"][row["ym"]]
            n = row["n"]
            cell(ws, r, 1, ru_month(row["ym"]), align=LEFT)
            cell(ws, r, 2, n)
            cell(ws, r, 3, SO_WINDOWS[f["zone"]][int(row["ym"][5:7])])
            cell(ws, r, 4, f"{row['best'][0]}–{row['best'][-1]}")
            for c, v in ((5, row["hit0"] / n), (6, row["hit1"] / n), (7, o["full"] / n)):
                cell(ws, r, c, round(v, 4), PCT, fill=OK if v >= ACCURACY else BAD)
            cell(ws, r, 8, round(o["safe"], 2), MW)
            cell(ws, r, 9, ", ".join(f"{h}: {k}" for h, k in row["modes"].most_common(4)),
                 align=LEFT)
            r += 1
        r += 1
    for col, w in {"A": 24, "B": 7, "C": 22, "D": 11, "E": 15, "F": 15, "G": 17,
                   "H": 17, "I": 26, "J": 12}.items():
        ws.column_dimensions[col].width = w


def build_method(ws):
    title(ws, "Методика", "Порядок расчёта")
    items = [
        ("ЧСМ", "Час собственного максимума субъекта — час наибольшего потребления "
                "региона в рабочие сутки, но только среди плановых пиковых часов "
                "СО ЕЭС (2-я ценовая зона: 5–8 и дневной интервал; 1-я: 8–21 с "
                "перерывами). Часы МСК, метки 1–24 как в отчётах АТС."),
        ("Разгрузка", "Завод снижает нагрузку на свой потолок в 4 часах подряд "
                      "внутри окна СО ЕЭС. Обязательство по мощности = нагрузка "
                      "завода в ЧСМ; если ЧСМ попал в блок, оно меньше на величину среза."),
        ("Сценарий 1", "Точность 80 %: зачётное снижение = потолок × 0,8. "
                       "Экономия в месяц = зачётное снижение × цена мощности года. "
                       "Окупаемость = инвестиции ÷ экономия."),
        ("Сценарий 2", "Срез завода снижает и нагрузку региона. Если срез больше "
                       "запаса пика блока над остальными часами окна, ЧСМ переезжает "
                       "в час без разгрузки и эффект суток = 0. Поэтому в дни "
                       "попадания срез ограничен этим запасом. Доля безопасного "
                       "среза = средний безопасный срез ÷ суммарный срез региона, "
                       "посчитана по посуточному факту АТС при идеальном выборе "
                       "блока. Заводы одного региона (Иркутская: БрАЗ, ИркАЗ, ТАЗ) "
                       "режут совместно, доля общая."),
        ("Данные", "Отчёты АТС «совокупное фактическое потребление в субъекте», "
                   "рабочие дни июнь 2025 — май 2026 (папка raw). Средний профиль "
                   "сверен с monthly_profiles. По Мурманской области посуточного "
                   "отчёта нет — только сценарий 1."),
        ("Не учтено", "Недовыпуск или перенос производства, разница цен "
                      "электроэнергии по часам, сетевая составляющая, реальные "
                      "инвестиционные затраты (принято 1 млн руб на завод)."),
    ]
    r = 4
    for h, body in items:
        c = ws.cell(r, 1, h)
        c.font, c.alignment = F_H, Alignment(vertical="top")
        c = ws.cell(r, 2, body)
        c.font = F
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 18 * (1 + len(body) // 95)
        r += 1
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 100


def main(argv):
    raw = Path(argv[1])
    out = Path(argv[2]) if len(argv) > 2 else Path(__file__).with_name(
        "Окупаемость_4ч_ЧСМ80.xlsx")
    data = load_ats.load_raw(str(raw))
    check_profiles(data, raw.parent if (raw.parent / "Иркутская область").is_dir() else None)
    facts = region_facts(data)
    for region, f in facts.items():
        print(f"{region:24} срез {f['depth']:5.1f} | фикс. блок {f['hit0']:.0%}/{f['hit1']:.0%}"
              f" | идеал: полный срез {f['full']:.0%}, безопасный {f['safe']:.1f} МВт "
              f"({f['safe'] / f['depth']:.0%})")

    wb = openpyxl.Workbook()
    wb.calculation.fullCalcOnLoad = True
    ws = wb.active
    ws.title = "1. Свод"
    build_svod(ws, facts)
    build_check(wb.create_sheet("2. Проверка по факту АТС"), facts)
    build_method(wb.create_sheet("3. Методика"))
    wb.save(out)
    print("Сохранено:", out)


if __name__ == "__main__":
    main(sys.argv)
